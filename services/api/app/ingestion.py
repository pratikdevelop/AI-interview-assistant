from __future__ import annotations

import base64
import io
import re
import zipfile
from dataclasses import asdict, dataclass
from pathlib import PurePosixPath
from urllib.parse import quote, urlparse

import httpx
from docx import Document
from pypdf import PdfReader

MAX_RESUME_BYTES = 8 * 1024 * 1024
MAX_ARCHIVE_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_FILES = 300
MAX_FILE_BYTES = 700 * 1024
MAX_TEXT_CHARS = 350_000
MAX_GITHUB_FILES = 35
ALLOWED_EXTENSIONS = {
    ".md", ".mdx", ".txt", ".py", ".js", ".jsx", ".ts", ".tsx", ".json",
    ".toml", ".yaml", ".yml", ".html", ".css", ".scss", ".sql", ".java",
    ".go", ".rs", ".php", ".rb", ".cs", ".sh", ".xml"
}
EXCLUDED_PARTS = {
    "node_modules", "vendor", "dist", "build", ".git", ".next", "coverage",
    "__pycache__", ".venv", "venv", "target", "Pods", ".idea", ".vscode"
}
SECRET_NAMES = {".env", "id_rsa", "id_ed25519", "credentials.json", "secrets.json"}


@dataclass
class Evidence:
    id: str
    source_kind: str
    source_label: str
    relative_path: str | None
    line_start: int | None
    line_end: int | None
    excerpt: str
    evidence_kind: str

    def to_dict(self) -> dict:
        return asdict(self)


def extract_resume(filename: str, content: bytes) -> str:
    if not content:
        raise ValueError("The resume file is empty.")
    if len(content) > MAX_RESUME_BYTES:
        raise ValueError("Resume is too large. The limit is 8 MB.")
    name = filename.lower()
    try:
        if name.endswith(".pdf") or content[:4] == b"%PDF":
            reader = PdfReader(io.BytesIO(content))
            if len(reader.pages) > 30:
                raise ValueError("Resume PDF exceeds the 30-page limit.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif name.endswith(".docx") or content[:2] == b"PK":
            doc = Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        elif name.endswith(".txt"):
            text = content.decode("utf-8", errors="replace")
        else:
            raise ValueError("Unsupported resume format. Upload a text-based PDF, DOCX, or TXT file.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not read this resume. Try a text-based PDF, DOCX, or TXT file.") from exc
    text = text.strip()
    if len(text) < 30:
        raise ValueError("Very little text could be extracted. Try a text-based PDF or DOCX.")
    return text[:MAX_TEXT_CHARS]


def _is_allowed_path(path: str) -> bool:
    p = PurePosixPath(path)
    low_parts = {part.lower() for part in p.parts}
    if p.is_absolute() or ".." in p.parts or low_parts.intersection({x.lower() for x in EXCLUDED_PARTS}):
        return False
    if any(part.startswith(".") and part.lower() not in {".github"} for part in p.parts[:-1]):
        return False
    filename = p.name.lower()
    if filename in SECRET_NAMES or filename.startswith(".env") or "secret" in filename or "credential" in filename:
        return False
    if p.suffix.lower() not in ALLOWED_EXTENSIONS:
        return False
    return True


def _make_evidence(path: str, raw_text: str, start_id: int, source_label: str = "Project source") -> list[Evidence]:
    # Keep excerpts small and line-addressable. Split long files into overlapping windows.
    lines = raw_text.replace("\x00", "").splitlines()
    output: list[Evidence] = []
    window = 45
    step = 38
    for start in range(0, min(len(lines), 450), step):
        chunk_lines = lines[start:start + window]
        chunk = "\n".join(chunk_lines).strip()
        if len(chunk) < 20:
            continue
        if len(chunk) > 5000:
            chunk = chunk[:5000]
        output.append(Evidence(
            id=f"ev_{start_id + len(output)}", source_kind="repository", source_label=source_label,
            relative_path=path, line_start=start + 1, line_end=start + len(chunk_lines),
            excerpt=chunk, evidence_kind="project_file"
        ))
        if len(output) >= 8:
            break
    return output


def ingest_zip(content: bytes) -> tuple[list[Evidence], str]:
    if not content:
        raise ValueError("The ZIP file is empty.")
    if len(content) > MAX_ARCHIVE_BYTES:
        raise ValueError("ZIP file exceeds the 20 MB upload limit.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise ValueError("The repository upload is not a valid ZIP archive.") from exc

    evidence: list[Evidence] = []
    total_expanded = 0
    file_count = 0
    next_id = 100
    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_FILES:
            raise ValueError(f"Archive contains too many entries; maximum is {MAX_FILES}.")
        for info in infos:
            if info.is_dir():
                continue
            file_count += 1
            total_expanded += info.file_size
            if total_expanded > MAX_EXPANDED_BYTES:
                raise ValueError("Archive expands beyond the 50 MB safety limit.")
            if file_count > MAX_FILES:
                raise ValueError("Archive has too many files.")
            path = info.filename.replace("\\", "/")
            if not _is_allowed_path(path) or info.file_size > MAX_FILE_BYTES:
                continue
            # Reject Unix symlinks from ZIP metadata.
            mode = (info.external_attr >> 16) & 0o170000
            if mode == 0o120000:
                continue
            if info.file_size == 0:
                continue
            try:
                raw = archive.read(info)
                text = raw.decode("utf-8")
            except (UnicodeDecodeError, RuntimeError, zipfile.BadZipFile):
                continue
            if not text.strip():
                continue
            evidence.extend(_make_evidence(path, text[:MAX_TEXT_CHARS], next_id, "Uploaded project"))
            next_id += 10
            if len(evidence) >= 100:
                break
    if not evidence:
        raise ValueError("No readable source or documentation files were found in this ZIP.")
    return evidence, f"Uploaded ZIP ({len(evidence)} evidence excerpts)"


async def ingest_github(url: str) -> tuple[list[Evidence], str]:
    parsed = urlparse(url.strip())
    if parsed.scheme != "https" or parsed.hostname not in {"github.com", "www.github.com"} or parsed.username or parsed.password:
        raise ValueError("Use a public repository URL in the form https://github.com/owner/repository.")
    parts = [p for p in parsed.path.strip("/").split("/") if p]
    if len(parts) != 2 or any(p in {".", ".."} for p in parts):
        raise ValueError("Enter the repository root URL, not a file, branch, or pull request URL.")
    owner, repo = parts
    repo = repo.removesuffix(".git")
    headers = {"Accept": "application/vnd.github+json", "User-Agent": "ProofPath-AI/0.1"}
    timeout = httpx.Timeout(12.0, connect=5.0)
    async with httpx.AsyncClient(timeout=timeout, follow_redirects=False, headers=headers) as client:
        meta_response = await client.get(f"https://api.github.com/repos/{quote(owner)}/{quote(repo)}")
        if meta_response.status_code == 404:
            raise ValueError("Repository not found or not public. Check the URL or upload a ZIP instead.")
        if meta_response.status_code == 403:
            raise ValueError("GitHub API rate limit reached. Please upload the repository as a ZIP instead.")
        if meta_response.status_code >= 400:
            raise ValueError("GitHub could not be reached. Upload the repository as a ZIP instead.")
        meta = meta_response.json()
        if meta.get("private"):
            raise ValueError("Private repositories are not supported in this MVP. Upload a safe ZIP instead.")
        branch = meta.get("default_branch") or "main"
        tree_response = await client.get(
            f"https://api.github.com/repos/{quote(owner)}/{quote(repo)}/git/trees/{quote(branch, safe='')}?recursive=1"
        )
        if tree_response.status_code == 403:
            raise ValueError("GitHub API rate limit reached. Upload a repository ZIP instead.")
        tree_response.raise_for_status()
        tree = tree_response.json().get("tree", [])
        chosen = [item for item in tree if item.get("type") == "blob" and _is_allowed_path(item.get("path", "")) and item.get("size", MAX_FILE_BYTES + 1) <= MAX_FILE_BYTES]
        # Prefer docs, manifests and app source; total number of files remains bounded.
        chosen.sort(key=lambda x: (0 if PurePosixPath(x["path"]).name.lower() in {"readme.md", "package.json", "pyproject.toml", "requirements.txt"} else 1, x["path"]))
        chosen = chosen[:MAX_GITHUB_FILES]
        evidence: list[Evidence] = []
        next_id = 100
        chars_used = 0
        for item in chosen:
            path = item["path"]
            contents_url = f"https://api.github.com/repos/{quote(owner)}/{quote(repo)}/contents/{quote(path, safe='/')}?ref={quote(branch, safe='')}"
            response = await client.get(contents_url)
            if response.status_code == 403:
                break
            if response.status_code >= 400:
                continue
            data = response.json()
            if data.get("encoding") != "base64" or not data.get("content"):
                continue
            try:
                raw = base64.b64decode(data["content"])
                text = raw.decode("utf-8")
            except (ValueError, UnicodeDecodeError):
                continue
            text = text[:min(MAX_TEXT_CHARS, MAX_TEXT_CHARS - chars_used)]
            if not text:
                break
            chars_used += len(text)
            evidence.extend(_make_evidence(path, text, next_id, "Public GitHub repository"))
            next_id += 10
            if len(evidence) >= 100 or chars_used >= MAX_TEXT_CHARS:
                break
    if not evidence:
        raise ValueError("No readable source or documentation files found. Try another repository or upload a ZIP.")
    return evidence, f"github.com/{owner}/{repo} ({len(evidence)} evidence excerpts)"


def evidence_to_dict(items: list[Evidence]) -> list[dict]:
    return [item.to_dict() for item in items]
